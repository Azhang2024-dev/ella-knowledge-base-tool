"""
md_converter.py — Markdown 转多格式核心库。
支持：PDF / DOCX / TXT / HTML
"""

import os
import re
import html as html_mod
from pathlib import Path

import markdown as md_lib
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


# ═══════════════════════════════════════════════════════════
# 工具函数
# ═══════════════════════════════════════════════════════════

def sanitize_xml(text: str) -> str:
    """移除 XML 不兼容的控制字符（保留 \\t \\n \\r）。"""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)


def strip_bom(text: str) -> str:
    """去掉文件开头的 BOM 字符。"""
    if text.startswith('﻿'):
        text = text[1:]
    return text


def clean_inline_md(text: str) -> str:
    """清理行内 Markdown 标记，保留纯文本。"""
    text = re.sub(r'<sup>.*?</sup>', '', text)
    text = re.sub(r'<sub>.*?</sub>', '', text)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'!\[.*?\]\(.*?\)', '', text)
    text = re.sub(r'\[([^\]]*)\]\(.*?\)', r'\1', text)
    text = re.sub(r'\$\$.*?\$\$', '[formula]', text, flags=re.DOTALL)
    text = re.sub(r'\$([^$]+)\$', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text


# ═══════════════════════════════════════════════════════════
# Markdown 表格解析器
# ═══════════════════════════════════════════════════════════

def parse_md_table(lines: list, start_idx: int):
    """
    从 start_idx 开始解析一个 Markdown 表格。
    返回 (headers, rows, end_idx) 或 None。
    """
    table_lines = []
    i = start_idx
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
        else:
            break

    if len(table_lines) < 2:
        return None

    def parse_row(line):
        cells = [c.strip() for c in line[1:-1].split('|')]
        return [clean_inline_md(c) for c in cells]

    sep_cells = parse_row(table_lines[1])
    is_separator = all(re.match(r'^:?-{3,}:?$', c) for c in sep_cells if c)

    if is_separator:
        headers = parse_row(table_lines[0])
        rows = [parse_row(line) for line in table_lines[2:]]
        return headers, rows, i
    else:
        rows = [parse_row(line) for line in table_lines]
        return [], rows, i


def table_to_txt(headers, rows):
    """将表格渲染为对齐的纯文本。"""
    if not rows:
        return ''
    all_rows = ([headers] if headers else []) + rows
    ncols = max(len(r) for r in all_rows)
    widths = [0] * ncols
    for r in all_rows:
        for ci in range(len(r)):
            widths[ci] = max(widths[ci], len(r[ci]))

    def fmt_row(row):
        padded = []
        for ci in range(ncols):
            c = row[ci] if ci < len(row) else ''
            padded.append(c.ljust(widths[ci]))
        return '| ' + ' | '.join(padded) + ' |'

    lines = []
    if headers:
        lines.append(fmt_row(headers))
        lines.append('|-' + '-|-'.join('-' * w for w in widths) + '-|')
    for r in rows:
        lines.append(fmt_row(r))
    return '\n'.join(lines)


def table_to_docx(doc, headers, rows):
    """在 docx 文档中插入原生表格。"""
    all_rows = ([headers] if headers else []) + rows
    if not all_rows:
        return
    ncols = max(len(r) for r in all_rows)
    nrows = len(all_rows)
    table = doc.add_table(rows=nrows, cols=ncols, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row_data in enumerate(all_rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            text = sanitize_xml(row_data[ci] if ci < len(row_data) else '')
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            if headers and ri == 0:
                run.bold = True
                shading = cell._element.get_or_add_tcPr()
                shading_elem = shading.makeelement(qn('w:shd'), {
                    qn('w:fill'): 'D9E2F3',
                    qn('w:val'): 'clear',
                })
                shading.append(shading_elem)
    doc.add_paragraph('')


# ═══════════════════════════════════════════════════════════
# MD → TXT
# ═══════════════════════════════════════════════════════════

def md_to_txt(text: str) -> str:
    """去除 Markdown 标记，保留纯文本（含格式化表格）。"""
    text = strip_bom(text)
    lines = text.split('\n')
    out_lines = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('|') and line.strip().endswith('|'):
            result = parse_md_table(lines, i)
            if result:
                headers, rows, end_i = result
                if rows:
                    out_lines.append('')
                    out_lines.append(table_to_txt(headers, rows))
                    out_lines.append('')
                i = end_i
                continue
            else:
                # parse_md_table 要求至少 2 行，单独 |...| 行当普通文本处理
                out_lines.append(clean_inline_md(line.strip()))
                i += 1
                continue
        h_match = re.match(r'^(#{1,6})\s+(.*)', line)
        if h_match:
            heading = clean_inline_md(h_match.group(2))
            if heading.strip():
                out_lines.append('')
                out_lines.append(heading.strip())
                out_lines.append('')
            i += 1
            continue
        if re.match(r'^[-\*]{3,}\s*$', line.strip()):
            out_lines.append('─' * 60)
            i += 1
            continue
        if not line.strip():
            out_lines.append('')
            i += 1
            continue
        if re.match(r'^!\[.*\]\(.*\)\s*$', line.strip()):
            i += 1
            continue
        list_match = re.match(r'^(\s*)[\*\-\+]\s+(.*)', line)
        if list_match:
            text_line = clean_inline_md(list_match.group(2))
            out_lines.append('• ' + text_line)
            i += 1
            continue
        ol_match = re.match(r'^(\s*)\d+\.\s+(.*)', line)
        if ol_match:
            text_line = clean_inline_md(ol_match.group(2))
            out_lines.append('  ' + text_line)
            i += 1
            continue
        if line.strip().startswith('>'):
            text_line = clean_inline_md(re.sub(r'^>\s?', '', line))
            out_lines.append('  ' + text_line)
            i += 1
            continue
        para_lines = []
        while i < len(lines) and lines[i].strip():
            l = lines[i]
            if re.match(r'^(#{1,6}\s|[-\*]{3,}\s*$)', l):
                break
            if l.strip().startswith('|') and l.strip().endswith('|'):
                break
            if re.match(r'^!\[.*\]\(.*\)\s*$', l.strip()):
                i += 1
                continue
            if re.match(r'^(\s*)[\*\-\+]\s+', l):
                break
            if re.match(r'^(\s*)\d+\.\s+', l):
                break
            if l.strip().startswith('>'):
                break
            para_lines.append(l)
            i += 1
        para_text = ' '.join(clean_inline_md(pl) for pl in para_lines)
        if para_text.strip():
            out_lines.append(para_text.strip())
    result = '\n'.join(out_lines)
    result = re.sub(r'\n{3,}', '\n\n', result)
    return result.strip()


# ═══════════════════════════════════════════════════════════
# MD → DOCX
# ═══════════════════════════════════════════════════════════

def md_to_docx(text: str, output_path: str):
    """将 Markdown 文本转为 .docx 文件（含原生表格）。"""
    text = strip_bom(text)
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('|') and line.strip().endswith('|'):
            result = parse_md_table(lines, i)
            if result:
                headers, rows, end_i = result
                if rows:
                    table_to_docx(doc, headers, rows)
                i = end_i
                continue
            else:
                # parse_md_table 要求至少 2 行，单独 |...| 行当普通段落处理
                p = doc.add_paragraph(sanitize_xml(clean_inline_md(line.strip())))
                p.paragraph_format.first_line_indent = Cm(0)
                i += 1
                continue
        h_match = re.match(r'^(#{1,6})\s+(.*)', line)
        if h_match:
            level = len(h_match.group(1))
            heading_text = clean_inline_md(h_match.group(2)).strip()
            if heading_text:
                heading_text = sanitize_xml(heading_text)
                if level <= 3:
                    doc.add_heading(heading_text, level=level)
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(heading_text)
                    run.bold = True
                    run.font.size = Pt(12)
            i += 1
            continue
        if re.match(r'^[-\*]{3,}\s*$', line.strip()):
            p = doc.add_paragraph()
            run = p.add_run('─' * 60)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            i += 1
            continue
        if not line.strip():
            doc.add_paragraph('')
            i += 1
            continue
        if re.match(r'^!\[.*\]\(.*\)\s*$', line.strip()):
            p = doc.add_paragraph()
            run = p.add_run('[Image]')
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            run.font.size = Pt(9)
            i += 1
            continue
        list_match = re.match(r'^(\s*)[\*\-\+]\s+(.*)', line)
        if list_match:
            text_line = sanitize_xml(clean_inline_md(list_match.group(2)))
            p = doc.add_paragraph()
            p.style = doc.styles['List Bullet']
            p.add_run(text_line)
            i += 1
            continue
        if line.strip().startswith('>'):
            text_line = sanitize_xml(clean_inline_md(re.sub(r'^>\s?', '', line)))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(text_line)
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue
        para_lines = []
        while i < len(lines) and lines[i].strip():
            l = lines[i]
            if re.match(r'^(#{1,6}\s|[-\*]{3,}\s*$)', l):
                break
            if l.strip().startswith('|') and l.strip().endswith('|'):
                break
            if re.match(r'^!\[.*\]\(.*\)\s*$', l.strip()):
                break
            if re.match(r'^(\s*)[\*\-\+]\s+', l):
                break
            if l.strip().startswith('>'):
                break
            para_lines.append(l)
            i += 1
        para_text = ' '.join(clean_inline_md(pl) for pl in para_lines)
        para_text = sanitize_xml(para_text)
        if para_text.strip():
            p = doc.add_paragraph(para_text.strip())
            p.paragraph_format.first_line_indent = Cm(0)
    doc.save(output_path)


# ═══════════════════════════════════════════════════════════
# MD → HTML
# ═══════════════════════════════════════════════════════════

def md_to_html(text: str, title: str = '') -> str:
    """将 Markdown 转为完整 HTML 文档。"""
    text = strip_bom(text)
    safe_title = html_mod.escape(title)
    body = md_lib.markdown(text, extensions=['tables', 'fenced_code', 'codehilite'])
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{safe_title}</title>
<style>
body {{ font-family: 'Times New Roman', serif; max-width: 800px; margin: 40px auto; line-height: 1.6; color: #222; }}
h1 {{ font-size: 1.5em; }} h2 {{ font-size: 1.3em; }} h3 {{ font-size: 1.15em; }} h4 {{ font-size: 1.05em; }}
table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
th, td {{ border: 1px solid #999; padding: 6px 10px; text-align: left; }}
th {{ background: #D9E2F3; font-weight: bold; }}
img {{ max-width: 100%; }}
blockquote {{ border-left: 3px solid #999; margin-left: 0; padding-left: 1em; color: #555; }}
sup {{ font-size: 0.75em; }}
pre {{ background: #f5f5f5; padding: 1em; overflow-x: auto; }}
code {{ background: #f0f0f0; padding: 1px 4px; border-radius: 3px; }}
</style>
</head>
<body>
{body}
</body>
</html>"""
    return html


# ═══════════════════════════════════════════════════════════
# DOCX → PDF（via Word COM）
# ═══════════════════════════════════════════════════════════

def open_word_app():
    """
    启动并返回一个 Word COM 实例，调用方负责最终调用 close_word_app()。
    返回 (word_app, pythoncom) 元组；失败时抛出异常。
    """
    import pythoncom
    pythoncom.CoInitialize()
    from win32com.client import Dispatch
    word = Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    return word, pythoncom


def close_word_app(word, pythoncom) -> None:
    """安全关闭 open_word_app() 返回的 Word 实例。"""
    try:
        if word is not None:
            word.Quit()
    except Exception:
        pass
    try:
        pythoncom.CoUninitialize()
    except Exception:
        pass


def docx_to_pdf_via_word(docx_path: str, pdf_path: str,
                         word_app=None) -> None:
    """
    将 DOCX 另存为 PDF。

    word_app: 已有的 Word COM 实例（open_word_app() 的返回值第一项）。
              传入时复用该实例，不再自行启停 Word，适合批量转换。
              不传时自行启停（兼容旧用法，但批量时有性能损耗）。
    """
    _own_word = word_app is None
    pythoncom_mod = None
    if _own_word:
        word_app, pythoncom_mod = open_word_app()
    try:
        doc = word_app.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close(SaveChanges=0)
    finally:
        if _own_word:
            close_word_app(word_app, pythoncom_mod)


# ═══════════════════════════════════════════════════════════
# 主转换入口
# ═══════════════════════════════════════════════════════════

def convert_file(md_path: str, target_format: str, output_dir: str,
                 temp_dir: str = None, word_app=None) -> str:
    """
    转换单个 .md 文件到目标格式。
    target_format: 'pdf' | 'docx' | 'txt' | 'html'
    word_app: 已有 Word COM 实例（仅 pdf 格式使用）；传入时批量复用，不传则单次启停。
    返回输出文件路径。
    """
    with open(md_path, 'r', encoding='utf-8') as f:
        text = f.read()

    stem = Path(md_path).stem
    output_dir = os.path.normpath(output_dir)

    if target_format == 'txt':
        out_path = os.path.normpath(os.path.join(output_dir, stem + '.txt'))
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(md_to_txt(text))
        return out_path

    elif target_format == 'docx':
        out_path = os.path.normpath(os.path.join(output_dir, stem + '.docx'))
        md_to_docx(text, out_path)
        return out_path

    elif target_format == 'pdf':
        import tempfile
        if temp_dir is None:
            temp_dir = output_dir
        temp_dir = os.path.normpath(temp_dir)
        # mkstemp 保证并发安全；suffix=.docx 让 Word 能识别格式
        fd, tmp_docx = tempfile.mkstemp(suffix='.docx', dir=temp_dir)
        os.close(fd)
        pdf_path = os.path.normpath(os.path.join(output_dir, stem + '.pdf'))
        try:
            md_to_docx(text, tmp_docx)
            docx_to_pdf_via_word(tmp_docx, pdf_path, word_app=word_app)
        finally:
            try:
                os.remove(tmp_docx)
            except Exception:
                pass
        return pdf_path

    elif target_format == 'html':
        out_path = os.path.normpath(os.path.join(output_dir, stem + '.html'))
        title = stem.replace('_', ' ')
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(md_to_html(text, title))
        return out_path

    else:
        raise ValueError(f'Unknown format: {target_format}')
